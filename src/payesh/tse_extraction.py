"""Best-effort TSE PDF terms and labeled Excel interest schedules."""

from datetime import date, datetime
from decimal import Decimal, InvalidOperation
import io
import json
import re
from pathlib import Path

import jdatetime
from openpyxl import load_workbook
import xlrd

from payesh.extraction_details import annotate
from payesh.extraction import _normalize, _date_value, extract_pdf_data
from payesh import web_storage as store


def number(value) -> str | None:
    if value is None or value == "":
        return None
    cleaned = _normalize(str(value)).replace(",", "").replace("٬", "").replace("،", "").strip()
    try:
        parsed = Decimal(cleaned)
        return format(parsed, "f") if parsed.is_finite() else None
    except InvalidOperation:
        return None


def date_value(value) -> str | None:
    if isinstance(value, (date, datetime)):
        gregorian = value.date() if isinstance(value, datetime) else value
        jalali = jdatetime.date.fromgregorian(date=gregorian)
        return f"{jalali.year:04d}/{jalali.month:02d}/{jalali.day:02d}"
    normalized = _normalize(str(value or ""))
    return _date_value(normalized) if re.fullmatch(r"\d{4}[/-]\d{1,2}[/-]\d{1,2}", normalized) else None


def _workbook_sheets(content: bytes, filename: str) -> list[tuple[str, list[tuple]]]:
    if Path(filename).suffix.lower() == ".xls":
        book = xlrd.open_workbook(file_contents=content)
        try:
            return [(sheet.name, [tuple(xlrd.xldate_as_datetime(cell.value, book.datemode)
                                       if cell.ctype == xlrd.XL_CELL_DATE else cell.value
                                       for cell in sheet.row(row)) for row in range(sheet.nrows)])
                    for sheet in book.sheets()]
        finally:
            book.release_resources()
    book = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        return [(sheet.title, list(sheet.iter_rows(values_only=True))) for sheet in book]
    finally:
        book.close()


def extract_excel(content: bytes, filename: str, symbol: str = "") -> tuple[dict, list[dict]]:
    """Read cached values; never calculate formulas or infer interest from unlabeled columns."""
    sheets = _workbook_sheets(content, filename)
    data = {"document_type": "daily_interest", "symbol": symbol or None,
            "sheets": [{"name": name, "rows": len(rows)} for name, rows in sheets]}
    daily: dict[str, str | None] = {}
    payments = []
    missing_values = 0
    recognized = False
    for name, rows in sheets:
        for row in rows[:100]:
            for column, value in enumerate(row[:-1]):
                label = _normalize(str(value or "")).replace("\u200c", " ")
                following = next((item for item in row[column + 1:] if item is not None and item != ""), None)
                if re.search(r"تاریخ\s*انتشار", label):
                    if parsed_date := date_value(following):
                        data["publication_date"] = parsed_date
                elif re.search(r"تاریخ\s*سررسید", label):
                    if parsed_date := date_value(following):
                        data["maturity_date"] = parsed_date
                elif re.fullmatch(r"\s*نرخ\s*سود\s*(?:سالانه)?\s*", label):
                    value_number = number(following)
                    if value_number is not None:
                        rate = Decimal(value_number)
                        data["annual_yield_percent"] = format(rate * 100 if abs(rate) <= 1 else rate, "f")
                elif re.search(r"(?:قیمت|مبلغ|ارزش)\s*اسمی", label):
                    if nominal := number(following):
                        data["nominal_value_rial"] = nominal
                elif "تعداد دفعات پرداخت سود در سال" in label:
                    count = number(following)
                    if count and Decimal(count) > 0:
                        months = Decimal(12) / Decimal(count)
                        if months == months.to_integral_value():
                            data["profit_payment_interval_months"] = str(int(months))
        for header_index, row in enumerate(rows[:30]):
            labels = [_normalize(str(value or "")).replace("\u200c", " ") for value in row]
            date_column = next((i for i, value in enumerate(labels) if value.strip() == "تاریخ"), None)
            amount_column = next((i for i, value in enumerate(labels)
                                  if "سود" in value and ("روزانه" in value or "روزشمار" in value or "روز شمار" in value)), None)
            if date_column is None or amount_column is None:
                continue
            recognized = True
            is_payment_schedule = "مواعد" in _normalize(name) or "پرداخت" in _normalize(name)
            for values in rows[header_index + 1:]:
                payment_date = date_value(values[date_column]) if date_column < len(values) else None
                if not payment_date:
                    continue
                amount = number(values[amount_column]) if amount_column < len(values) else None
                missing_values += amount is None
                if is_payment_schedule:
                    payments.append({"payment_date": payment_date, "amount_rial": amount})
                else:
                    daily[payment_date] = amount
            break
    if payments:
        data["payment_schedule"] = payments
    required = ("publication_date", "maturity_date", "annual_yield_percent", "nominal_value_rial")
    data["extraction_status"] = ("unrecognized" if not recognized else
                                 "complete" if daily and not missing_values and all(data.get(key) is not None for key in required) else "partial")
    if missing_values:
        data["missing_cached_values"] = missing_values
    data["fields_found"] = [key for key, value in data.items() if value is not None]
    values = [{"payment_date": key, "amount_rial": value} for key, value in sorted(daily.items())]
    annotate(data, values, required)
    return data, values


def extract_tse_attachments(path: Path) -> tuple[int, int, int, int]:
    with store.connection(path) as db:
        records = [dict(row) for row in db.execute("""SELECT f.*, a.symbol, a.title, a.abstract
            FROM source_attachments f JOIN announcements a USING(notice_number)
            WHERE a.source='tse' AND f.file_blob IS NOT NULL ORDER BY f.kind DESC, f.attachment_id""")]
        notices = [row[0] for row in db.execute("SELECT notice_number FROM announcements WHERE source='tse'")]
    for record in records:
        try:
            if record["kind"] == "excel":
                data, values = extract_excel(record["file_blob"], record["filename"], record["symbol"])
            else:
                data, values = extract_pdf_data(record["file_blob"])
                if not data.get("symbol") and record["symbol"]:
                    data["symbol"] = record["symbol"]
            data["daily_interest_values"] = values
        except Exception as exc:
            data = {"document_type": "unknown", "extraction_status": "error", "error": str(exc), "daily_interest_values": []}
        if record["kind"] == "pdf" or data.get("extraction_status") == "error":
            annotate(data, data.get("daily_interest_values", []))
        with store.connection(path) as db:
            db.execute("UPDATE source_attachments SET extracted_data_json=?,extraction_status=? WHERE attachment_id=?",
                       (json.dumps(data, ensure_ascii=False), data["extraction_status"], record["attachment_id"]))
    complete = partial = daily_count = 0
    for notice in notices:
        with store.connection(path) as db:
            attachments = db.execute("""SELECT kind,filename,extracted_data_json,extraction_status FROM source_attachments
                WHERE notice_number=? ORDER BY CASE WHEN kind='pdf' THEN 0 ELSE 1 END,attachment_id""", (notice,)).fetchall()
            parsed = [(row["kind"], json.loads(row["extracted_data_json"])) for row in attachments if row["extracted_data_json"]]
            if not parsed:
                continue
            # Prefer terms from a parsed PDF, enriching missing terms from the workbook.
            merged: dict = {}
            daily_values = []
            for kind, data in parsed:
                for key, value in data.items():
                    if key not in {"extraction_status", "daily_interest_values", "fields_found", "document_type", "error", "missing_fields", "parsing_reasons", "missing_cached_values"} and value not in (None, "", []):
                        if key not in merged:
                            merged[key] = value
                if data.get("daily_interest_values") and (kind == "excel" or not daily_values):
                    daily_values = data["daily_interest_values"]
            statuses = [data["extraction_status"] for _, data in parsed]
            all_downloaded = all(row["extracted_data_json"] for row in attachments)
            status = "complete" if "complete" in statuses and all_downloaded and "error" not in statuses else "partial"
            if all(value == "error" for value in statuses):
                status = "error"
            elif all(value in {"unrecognized", "no_text"} for value in statuses):
                status = "unrecognized"
            # The title determines the notice type; a daily-interest attachment is supporting data.
            title = db.execute("SELECT title FROM announcements WHERE notice_number=?", (notice,)).fetchone()[0]
            document_type = "opening" if "گشایش" in title else "subscription"
            merged.update(document_type=document_type, extraction_status=status)
            annotate(merged, daily_values)
            merged["parsing_reasons"].extend({"code": "missing_file", "filename": row["filename"]} for row in attachments if not row["extracted_data_json"])
            if "error" in statuses and status != "error":
                merged["parsing_reasons"].append({"code": "attachment_error"})
            merged["fields_found"] = [key for key in merged if key not in {"missing_fields", "parsing_reasons"}]
            db.execute("""INSERT INTO announcement_extractions VALUES(?,?,?,?,?)
                ON CONFLICT(notice_number) DO UPDATE SET document_type=excluded.document_type,
                extraction_status=excluded.extraction_status,extracted_data_json=excluded.extracted_data_json,
                extracted_at=excluded.extracted_at""",
                       (notice, document_type, status, json.dumps(merged, ensure_ascii=False), store.now()))
            db.execute("DELETE FROM daily_interest_values WHERE notice_number=?", (notice,))
            db.executemany("INSERT INTO daily_interest_values VALUES(?,?,?)",
                           [(notice, row["payment_date"], row["amount_rial"]) for row in daily_values])
            complete += status == "complete"
            partial += status != "complete"
            daily_count += len(daily_values)
    return len(records), complete, partial, daily_count
