"""TSE matching, first-page scope, attachment extraction, migration, and source isolation."""

from contextlib import contextmanager
import io
import json
import sqlite3
from urllib.parse import parse_qs, urlsplit

from fastapi.testclient import TestClient
from openpyxl import Workbook
import pytest

from payesh import service, tse, web_storage as store
from payesh.database import _SCHEMA, save_announcements, save_announcement_file
from payesh.ifb import Announcement
from payesh.tse_extraction import extract_excel, extract_tse_attachments
from payesh.web import create_app


def payload(title="اطلاعیه پذیره‌نویسی اوراق مرابحه", source_id=123):
    return {"items": [{"title": title, "abstract": "با نماد معاملاتی «صدشت907»",
                       "content": {"id": source_id, "publishAt": "2026-10-07T17:07:00", "files": [
                           {"fileId": "pdf-id", "fileName": "anything.pdf", "link": "https://webgw.tse.ir/File/Download/pdf-id"},
                           {"fileId": "excel-id", "fileName": "schedule.xlsx", "link": "https://webgw.tse.ir/File/Download/excel-id"},
                           {"fileId": "image", "fileName": "banner.jpg", "link": "https://webgw.tse.ir/File/Download/image", "isMasterImage": True}]} }],
            "hasNextPage": True, "totalPages": 1232, "totalCount": 24640}


def workbook_bytes(cached_amount=0.0):
    book = Workbook()
    sheet = book.active
    sheet.title = "جدول سود روزانه"
    sheet.append(["تاریخ", "سود روزانه تجمعی هر ورقه", None, "تاريخ انتشار", "1405/07/19"])
    sheet.append(["1405/07/19", cached_amount, None, "تاريخ سررسيد", "1409/07/19"])
    sheet.append(["1405/07/20", 593.9041095890411, None, "نرخ سود", 0.23])
    sheet.append([None, None, None, "قیمت اسمی (ریال)", 1000000])
    sheet.append([None, None, None, "تعداد دفعات پرداخت سود در سال", 4])
    book.create_sheet("جدول مواعد پرداخت").append(["تاریخ", "سود روزانه تجمعی هر ورقه"])
    book["جدول مواعد پرداخت"].append(["1405/10/19", 56712.32876712329])
    stream = io.BytesIO()
    book.save(stream)
    return stream.getvalue()


@pytest.mark.parametrize("title, expected", [
    ("پذیره نویسی اوراق", True), ("اوراق مرابحه: گشایش", True),
    ("اطلاعيه پذیره\u200cنویسی اوراق", True), ("پذیره نویسی صندوق سهام", False),
    ("گشایش بازار سهام", False), ("معاملات اوراق", False), ("پذیره اور\u200cاق", True),
])
def test_title_rule(title, expected):
    assert tse.matches_title(title) is expected


def test_fetch_requests_only_first_page(monkeypatch):
    requested = []
    @contextmanager
    def fake_urlopen(request, timeout):
        requested.append(request.full_url)
        yield io.BytesIO(json.dumps(payload()).encode())
    monkeypatch.setattr(tse.urllib.request, "urlopen", fake_urlopen)
    records = tse.fetch_announcements()
    assert len(requested) == 1
    query = parse_qs(urlsplit(requested[0]).query)
    assert query["PageNumber"] == ["1"] and query["PageSize"] == ["20"]
    assert records[0].source_id == "123"
    assert records[0].published_at == "17:07:00 1405/07/15"
    assert records[0].symbol == "صدشت907"
    assert len(records[0].attachments) == 2


def test_legacy_migration_preserves_ifb_data(tmp_path):
    path = tmp_path / "legacy.sqlite3"
    with sqlite3.connect(path) as db:
        db.executescript(_SCHEMA)
        db.execute("INSERT INTO announcements VALUES(?,?,?,?,?,?,?,?)",
                   ("123", "1", "IFB", "old title", "1405/07/15", "https://ifb.ir/file", "first", "last"))
        db.execute("INSERT INTO announcement_files VALUES(?,?,?,?,?,?)",
                   (1, "123", "old.pdf", "application/pdf", b"%PDF-old", "downloaded"))
        db.execute("INSERT INTO announcement_extractions VALUES(?,?,?,?,?)",
                   ("123", "subscription", "partial", '{"symbol":"IFB"}', "extracted"))
    store.initialize(path)
    tse.save_announcements(tse.parse_announcements(payload()), path)
    store.initialize(path)
    old = store.detail(path, "123")
    assert old["source"] == "ifb" and old["source_id"] == "123"
    assert old["first_seen_at"] == "first"
    assert old["extracted_data"]["symbol"] == "IFB"
    assert store.pdf(path, "123")[1] == b"%PDF-old"
    assert store.announcements(path)["total"] == 2
    assert store.detail(path, "tse:123")["source"] == "tse"


def test_multiple_attachments_retry_and_source_filter(tmp_path):
    path = tmp_path / "files.sqlite3"
    store.initialize(path)
    records = tse.parse_announcements(payload())
    tse.save_announcements(records, path)
    tse.save_announcements(records, path)
    assert len(tse.missing_attachments(path)) == 2
    excel = next(file for file in tse.missing_attachments(path) if file["kind"] == "excel")
    tse.save_attachment(path, excel["attachment_id"], workbook_bytes(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    assert len(tse.missing_attachments(path)) == 1
    with TestClient(create_app(path, update_interval=0)) as client:
        assert client.get("/api/v1/announcements", params={"source": "tse", "date_from": "1405-07-15", "date_to": "1405-07-15"}).json()["total"] == 1
        assert client.get("/api/v1/announcements", params={"source": "ifb"}).json()["total"] == 0
        detail = client.get("/api/v1/announcements/tse:123").json()
        assert len(detail["attachments"]) == 2
        response = client.get(f"/api/v1/announcements/tse:123/attachments/{excel['attachment_id']}")
        assert response.content.startswith(b"PK")
        assert response.headers["content-disposition"].startswith("attachment;")
        assert client.get(f"/api/v1/announcements/wrong/attachments/{excel['attachment_id']}").status_code == 404
    save_announcements([], path)
    assert store.detail(path, "tse:123") is not None


def test_excel_labeled_terms_zero_values_and_schedules():
    data, values = extract_excel(workbook_bytes(), "unrestricted-name.xlsx", "صدشت907")
    assert data["extraction_status"] == "complete"
    assert data["annual_yield_percent"] == "23.00"
    assert data["nominal_value_rial"] == "1000000"
    assert data["profit_payment_interval_months"] == "3"
    assert values[0]["amount_rial"] == "0"
    assert values[1]["amount_rial"] == "593.9041095890411"
    assert len(values) == 2
    assert len(data["payment_schedule"]) == 1
    incomplete, _ = extract_excel(workbook_bytes(cached_amount="=A2+1"), "formulas.xlsx")
    assert incomplete["extraction_status"] == "partial"
    assert incomplete["missing_cached_values"] == 1


def test_unlabeled_workbook_is_not_guessed():
    book = Workbook()
    book.active.append(["1405/07/19", 999])
    stream = io.BytesIO(); book.save(stream)
    data, rows = extract_excel(stream.getvalue(), "unknown.xlsx")
    assert data["extraction_status"] == "unrecognized" and rows == []


def test_tse_extraction_without_filename_rule_and_missing_attachment_status(tmp_path):
    path = tmp_path / "extract.sqlite3"
    store.initialize(path)
    tse.save_announcements(tse.parse_announcements(payload()), path)
    excel = next(file for file in tse.missing_attachments(path) if file["kind"] == "excel")
    tse.save_attachment(path, excel["attachment_id"], workbook_bytes(), "application/octet-stream")
    extract_tse_attachments(path)
    detail = store.detail(path, "tse:123")
    assert detail["extraction_status"] == "partial"  # PDF is still pending.
    assert len(detail["daily_interest_values"]) == 2
    assert detail["extracted_data"]["nominal_value_rial"] == "1000000"


def test_source_failures_are_independent_and_downloads_are_idempotent(tmp_path, monkeypatch):
    path = tmp_path / "sources.sqlite3"
    def failed_ifb():
        raise RuntimeError("IFB offline")
    monkeypatch.setattr(service, "fetch_announcements", failed_ifb)
    monkeypatch.setattr(tse, "fetch_announcements", lambda: tse.parse_announcements(payload()))
    downloads = []
    def fake_download(file):
        downloads.append(file["source_file_id"])
        return (workbook_bytes(), "application/octet-stream") if file["kind"] == "excel" else (b"%PDF-broken", "application/pdf")
    monkeypatch.setattr(tse, "download_attachment", fake_download)
    first = service.collect(path)
    assert first["sources"]["ifb"]["status"] == "failed"
    assert first["sources"]["tse"]["fetched"] == 1
    assert first["downloaded"] == 2
    second = service.collect(path)
    assert second["downloaded"] == 0 and len(downloads) == 2
    assert store.announcements(path)["total"] == 1


def test_attachment_download_rejects_wrong_host_and_wrong_file(monkeypatch):
    with pytest.raises(ValueError, match="host"):
        tse.download_attachment({"file_url": "https://example.com/file", "kind": "pdf"})
    class FakeResponse(io.BytesIO):
        def geturl(self):
            return "https://webgw.tse.ir/File/Download/file"
    monkeypatch.setattr(tse.urllib.request, "urlopen", lambda request, timeout: FakeResponse(b"<html>error</html>"))
    with pytest.raises(ValueError, match="PDF"):
        tse.download_attachment({"file_url": "https://webgw.tse.ir/File/Download/file", "kind": "pdf"})
